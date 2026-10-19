"""
Exporta para Excel os dados dos painéis Power BI públicos da ABVE.

Os relatórios foram publicados com "Publicar na web", e o visualizador
público do Power BI consulta o modelo semântico por uma API anônima
(`/public/reports/querydata`), autorizada apenas pela chave do link.
Este script faz as mesmas consultas que o navegador faria:

  1. lê o esquema conceitual do modelo (tabelas e colunas visíveis);
  2. pede todas as linhas de cada tabela, paginando com restart tokens;
  3. grava um .xlsx por painel, com uma aba por tabela e uma aba "Sobre".

Medidas DAX não são exportadas (dependem de contexto de filtro); as
colunas trazem o dado de base de onde elas são calculadas.

A API sempre agrupa: pedir só colunas devolveria combinações distintas,
e linhas idênticas da tabela de fatos (cada emplacamento vale QTD=1)
seriam fundidas. Por isso as colunas de quantidade vão como SOMA,
agrupadas pelas demais — o total bate com o do modelo.

Uso:  python etl/export_powerbi_abve.py
Saída: data/powerbi_abve/<painel>.xlsx

Dependência: openpyxl  (pip install openpyxl)
"""
import base64
import datetime as dt
import gzip
import json
import re
import sys
import time
import urllib.request
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
SAIDA = RAIZ / "data" / "powerbi_abve"

CLUSTER = "https://wabi-brazil-south-b-primary-api.analysis.windows.net/public/reports"
POR_PAGINA = 30000

PAINEIS = [
    ("Eletromobilidade", "https://abve.org.br/abve-data/bi-geral/",
     "eyJrIjoiZDczMmRiZGEtMGVlMy00NTg2LThiNzUtOWQxZjIyZjNhOWFhIiwidCI6IjQwNmMyNTc1LWM5ZDQtNGQ0NC1hNWI2LWJiNThiYzg2MDMzZSJ9"),
    ("Frotas", "https://abve.org.br/abve-data/bi-frotas/",
     "eyJrIjoiMDc5ZWMzYjgtMWZkNi00YzAwLTkzYjAtZmQyMmI5MjU4OTcwIiwidCI6IjQwNmMyNTc1LWM5ZDQtNGQ0NC1hNWI2LWJiNThiYzg2MDMzZSJ9"),
    ("Geografia da Eletromobilidade", "https://abve.org.br/abve-data/bi-geografia-da-eletromobilidade/",
     "eyJrIjoiMzg5MzllM2UtMTJlNC00YTZkLThiODctYTdiMTYzZDE3YWVkIiwidCI6IjQwNmMyNTc1LWM5ZDQtNGQ0NC1hNWI2LWJiNThiYzg2MDMzZSJ9"),
    ("Ônibus", "https://abve.org.br/abve-data/bi-onibus/",
     "eyJrIjoiZThmN2FmYmQtYzZhNy00ZDk3LTkzNDYtYjNmZWJkMjVlMGQyIiwidCI6IjQwNmMyNTc1LWM5ZDQtNGQ0NC1hNWI2LWJiNThiYzg2MDMzZSJ9"),
    ("Eletropostos", "https://abve.org.br/abve-data/bi-eletropostos/",
     "eyJrIjoiYzRiNjM5M2EtMzE0NS00ZjQ2LTljZTMtODVlZjc4YWQ2M2Y1IiwidCI6IjQwNmMyNTc1LWM5ZDQtNGQ0NC1hNWI2LWJiNThiYzg2MDMzZSJ9"),
    ("Eletrificados+MHEV", "https://abve.org.br/abve-data/eletrificadosmhev/",
     "eyJrIjoiMzAxYzQ4ZjUtMjkwOC00ZDhlLTgwNWMtZDkwOGFlMmMyZDdmIiwidCI6IjQwNmMyNTc1LWM5ZDQtNGQ0NC1hNWI2LWJiNThiYzg2MDMzZSJ9"),
]

# tabelas de data geradas automaticamente pelo Power BI: só repetem datas
IGNORAR = re.compile(r"^(LocalDateTable_|DateTableTemplate_)")

# numéricas com resumo padrão "Soma" que NÃO são quantidades
NAO_SOMAR = re.compile(r"codigo|ordem|ano|m[eê]s|latitude|longitude|n[uú]mero", re.I)
TIPOS_NUM = (3, 4)  # Double, Int64


def somavel(p):
    return (p.get("DataType") in TIPOS_NUM
            and p["Column"].get("DefaultAggregate") == 2
            and not NAO_SOMAR.search(p["Name"]))


def http(url, chave, corpo=None):
    dados = json.dumps(corpo).encode() if corpo is not None else None
    req = urllib.request.Request(url, data=dados, headers={
        "X-PowerBI-ResourceKey": chave,
        "Content-Type": "application/json",
        "Accept-Encoding": "gzip",
        "User-Agent": "Mozilla/5.0 (SustAmbiTech ETL)",
    })
    for tentativa in range(4):
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                bruto = r.read()
                if r.headers.get("Content-Encoding") == "gzip":
                    bruto = gzip.decompress(bruto)
                return json.loads(bruto.decode("utf-8"))
        except Exception as e:  # noqa: BLE001 - rede instável: tenta de novo
            if tentativa == 3:
                raise
            print(f"    ! {e} — nova tentativa", file=sys.stderr)
            time.sleep(3 * (tentativa + 1))


def chave_do_link(r):
    pad = "=" * (-len(r) % 4)
    return json.loads(base64.b64decode(r + pad))["k"]


# ---------------------------------------------------------------- DSR
def decodificar(resposta, nomes):
    """Converte o formato DSR do Power BI em lista de linhas.

    As colunas agrupadas (G0, G1…) e as agregações (M0…) vêm em blocos
    separados, não na ordem do Select; o descritor da resposta diz qual
    chave corresponde a qual nome, e é por ele que as linhas são montadas.
    """
    res = resposta["results"][0]["result"]["data"]
    chave_de = {s["Name"]: s["Value"] for s in res["descriptor"]["Select"]}
    ds = res["dsr"]["DS"][0]
    dicts = ds.get("ValueDicts", {})
    linhas_brutas = ds["PH"][0].get("DM0", []) if ds.get("PH") else []
    esquema, anterior, linhas = None, None, []
    for lin in linhas_brutas:
        if "S" in lin:
            esquema = lin["S"]
        repetir = lin.get("R", 0)
        nulos = lin.get("Ø", 0)
        valores = iter(lin.get("C", []))
        atual = {}
        for i, col in enumerate(esquema):
            n = col["N"]
            if repetir >> i & 1:
                atual[n] = anterior[n]
            elif nulos >> i & 1:
                atual[n] = None
            else:
                v = next(valores)
                if col.get("DN") and isinstance(v, int):
                    v = dicts[col["DN"]][v]
                atual[n] = v
        linhas.append(atual)
        anterior = atual
    tipos = {c["N"]: c.get("T") for c in (esquema or [])}
    chaves = [chave_de.get(n) for n in nomes]
    return ([[l.get(k) for k in chaves] for l in linhas],
            [tipos.get(k) for k in chaves], ds.get("RT"))


def converter(v, tipo):
    if v is None:
        return None
    if tipo == 7 and isinstance(v, (int, float)):           # DateTime em epoch ms
        return dt.datetime(1970, 1, 1) + dt.timedelta(milliseconds=v)
    if isinstance(v, str) and tipo in (3, 4) and re.fullmatch(r"-?\d+(\.\d+)?", v):  # números em texto
        return float(v) if "." in v else int(v)
    return v


def consultar_tabela(chave, modelo, entidade, colunas, somar=()):
    def col(c):
        return {"Column": {"Expression": {"SourceRef": {"Source": "t"}}, "Property": c}}

    sel = [{"Aggregation": {"Expression": col(c), "Function": 0}, "Name": f"Sum(t.{c})"}
           if c in somar else dict(col(c), Name=f"t.{c}") for c in colunas]
    tokens, todas, tipos = None, [], None
    while True:
        janela = {"Count": POR_PAGINA}
        if tokens:
            janela["RestartTokens"] = tokens
        corpo = {
            "version": "1.0.0",
            "queries": [{
                "Query": {"Commands": [{"SemanticQueryDataShapeCommand": {
                    "Query": {"Version": 2,
                              "From": [{"Name": "t", "Entity": entidade, "Type": 0}],
                              "Select": sel},
                    "Binding": {
                        "Primary": {"Groupings": [{"Projections": list(range(len(sel)))}]},
                        "DataReduction": {"DataVolume": 6, "Primary": {"Window": janela}},
                        "Version": 1},
                    "ExecutionMetricsKind": 1}}]},
                "QueryId": "",
                "ApplicationContext": {"DatasetId": modelo["dbName"],
                                       "Sources": [{"ReportId": modelo["report"]}]},
            }],
            "cancelQueries": [],
            "modelId": modelo["id"],
        }
        resp = http(f"{CLUSTER}/querydata?synchronous=true", chave, corpo)
        erro = resp["results"][0]["result"].get("data", {}).get("dsr", {}).get("DataShapes")
        if erro:
            raise RuntimeError(json.dumps(erro)[:300])
        linhas, t, novos = decodificar(resp, [x["Name"] for x in sel])
        tipos = tipos or t
        todas.extend(linhas)
        if not novos or not linhas:
            break
        tokens = novos
        time.sleep(0.5)
    tipos = tipos or [None] * len(colunas)
    return [[converter(v, tipos[i]) for i, v in enumerate(l)] for l in todas]


# ---------------------------------------------------------------- Excel
CAB = Font(bold=True, color="FFFFFF")
FUNDO = PatternFill("solid", fgColor="2F4A44")


def nome_aba(nome, usados):
    base = re.sub(r"[\[\]\*\?/\\:]", "_", nome)[:31] or "Tabela"
    n, i = base, 2
    while n.lower() in usados:
        sufixo = f"_{i}"
        n, i = base[:31 - len(sufixo)] + sufixo, i + 1
    usados.add(n.lower())
    return n


def gravar_aba(wb, titulo, colunas, linhas):
    ws = wb.create_sheet(titulo)
    ws.append(colunas)
    for c in ws[1]:
        c.font, c.fill = CAB, FUNDO
    for l in linhas:
        ws.append(l)
    ws.freeze_panes = "A2"
    if linhas:
        ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(colunas, 1):
        amostra = [len(str(col))] + [len(str(l[i - 1])) for l in linhas[:200] if l[i - 1] is not None]
        ws.column_dimensions[get_column_letter(i)].width = min(max(amostra) + 2, 60)
        if any(isinstance(l[i - 1], dt.datetime) for l in linhas[:50]):
            for cel in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cel[0].number_format = "dd/mm/yyyy"


def exportar(nome, origem, r):
    chave = chave_do_link(r)
    print(f"\n== {nome}")
    me = http(f"{CLUSTER}/{chave}/modelsAndExploration?preferReadOnlySession=true", chave)
    m = me["models"][0]
    modelo = {"id": m["id"], "dbName": m["dbName"],
              "report": me["exploration"]["report"]["objectId"]}
    cs = http(f"{CLUSTER}/conceptualschema", chave,
              {"modelIds": [m["id"]], "userPreferredLocale": "pt-BR"})
    entidades = cs["schemas"][0]["schema"]["Entities"]

    wb = Workbook()
    sobre = wb.active
    sobre.title = "Sobre"
    usados = {"sobre"}
    resumo = []
    for e in entidades:
        if e.get("Hidden") or IGNORAR.match(e["Name"]):
            continue
        props = [p for p in e["Properties"] if "Column" in p and not p.get("Hidden")]
        colunas = [p["Name"] for p in props]
        somar = {p["Name"] for p in props if somavel(p)}
        if not colunas:
            continue
        if somar == set(colunas):   # só quantidades: nada para agrupar
            somar = set()
        try:
            linhas = consultar_tabela(chave, modelo, e["Name"], colunas, somar)
        except Exception as erro:  # noqa: BLE001 - uma tabela ruim não derruba o painel
            print(f"  x {e['Name']}: {erro}")
            resumo.append((e["Name"], "erro", str(erro)[:120]))
            continue
        aba = nome_aba(e["Name"], usados)
        gravar_aba(wb, aba, colunas, linhas)
        resumo.append((e["Name"], len(linhas), aba))
        print(f"  {len(linhas):>7} linhas  {e['Name']}")
        time.sleep(0.5)

    sobre.append([f"Painel ABVE — {nome}"])
    sobre["A1"].font = Font(bold=True, size=14)
    sobre.append(["Página da ABVE", origem])
    sobre.append(["Relatório Power BI", f"https://app.powerbi.com/view?r={r}"])
    sobre.append(["Exportado em", dt.datetime.now().strftime("%d/%m/%Y %H:%M")])
    sobre.append(["Autoria dos dados", "ABVE — Associação Brasileira do Veículo Elétrico"])
    sobre.append(["Observação", "Colunas do modelo semântico publicado; medidas DAX não incluídas."])
    sobre.append([])
    sobre.append(["Tabela do modelo", "Linhas", "Aba"])
    for c in sobre[sobre.max_row]:
        c.font, c.fill = CAB, FUNDO
    for linha in resumo:
        sobre.append(list(linha))
    sobre.column_dimensions["A"].width = 48
    sobre.column_dimensions["B"].width = 70
    sobre.column_dimensions["C"].width = 34

    SAIDA.mkdir(parents=True, exist_ok=True)
    arquivo = SAIDA / (re.sub(r"[^\w\-]+", "_", nome).strip("_") + ".xlsx")
    wb.save(arquivo)
    print(f"  -> {arquivo.relative_to(RAIZ)}")


if __name__ == "__main__":
    for painel in PAINEIS:
        exportar(*painel)
